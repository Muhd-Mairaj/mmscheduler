import openpyxl
import json
import my_sheet


def main():
    # change this to the path of the excel file
    excel_file_path = "timetable_sem3.xlsx"
    json_file_path = "one_week_schedule.json"
    json_file_path_occ_separated = "one_week_schedule_occ_separated.json"
    workbook = openpyxl.load_workbook(excel_file_path, read_only=True)
    sheet = my_sheet.get_sheet(workbook.active)

    data = read_data(sheet)
    json.dump(data, open(json_file_path, "w"), indent=2)

    data_occ_separated = read_data_occ_separated(sheet)
    json.dump(data_occ_separated, open(json_file_path_occ_separated, "w"), indent=2)


def read_data_occ_separated(sheet: my_sheet.MyReadOnlyWorksheet | my_sheet.Worksheet):
    data = []
    tracker = {}
    for row in sheet.iter_rows(min_row=sheet.header_row+1, max_row=sheet.end_row, values_only=True):
        module_offering = row[sheet._module_offering_column - 1]
        code, occurence = parse_module_offering(module_offering)

        module = row[sheet._module_column - 1]
        activity = row[sheet._activity_column - 1]
        room = row[sheet._room_column - 1]
        day = row[sheet._day_column - 1]
        begin_time = row[sheet._begin_column - 1]
        end_time = row[sheet._end_column - 1]

        for occ in occurence:
            if code not in tracker:
                tracker[code] = {}
            
            # if this occ for this code is not found before, track it
            if occ not in tracker[code]:
                tracker[code][occ] = {
                    "module": module,
                    "course_id": code,
                    "occurence": occ,
                    activity.lower(): {
                        "day": day,
                        "room": room,
                        "begin_time": begin_time,
                        "end_time": end_time,
                    } 
                }
            # otherwise, update it and add to data
            else:
                tracker[code][occ][activity.lower()] = {
                    "day": day,
                        "room": room,
                        "begin_time": begin_time,
                        "end_time": end_time,
                }
                data.append(tracker[code][occ])

    return data


def read_data(sheet: my_sheet.MyReadOnlyWorksheet | my_sheet.Worksheet):
    data = {}
    for row in sheet.iter_rows(min_row=sheet.header_row+1, max_row=sheet.end_row, values_only=True):
        module_offering = row[sheet._module_offering_column - 1]
        code, occurence = parse_module_offering(module_offering)

        module = row[sheet._module_column - 1]
        activity = row[sheet._activity_column - 1]
        room = row[sheet._room_column - 1]
        day = row[sheet._day_column - 1]
        begin_time = row[sheet._begin_column - 1]
        end_time = row[sheet._end_column - 1]

        if code not in data:
            data[code] = {}
            data[code]["Module"] = module

        if day not in data[code]:
            data[code][day] = []

        data[code][day].append({
            "Occurences": occurence,
            "Activity": activity,
            "Room": room,
            "Begin Time": begin_time,
            "End Time": end_time,
        })

    return data


def parse_module_offering(module_offering: str) -> tuple:
    offerings = module_offering.split(", ")
    if len(offerings) != 1:
        code = offerings[0].split("/")[0]
        occ = [offering.split("/")[-1] for offering in offerings]
        return code, occ

    parsed = module_offering.split("/")
    return parsed[0], [parsed[-1]]


if __name__ == "__main__":
    main()
